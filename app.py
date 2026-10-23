from __future__ import annotations

import io
import os
import re
import shutil
import sqlite3
import sys
import threading
import unicodedata
import webbrowser
import zipfile
from contextlib import contextmanager
from datetime import date
from email.message import EmailMessage
from email.policy import SMTP
from pathlib import Path
from typing import Any, Iterator

import openpyxl
from docx import Document
from docx.opc.exceptions import PackageNotFoundError
from docx.shared import Pt
from docx.oxml.ns import qn
from flask import Flask, jsonify, request, send_file
from werkzeug.serving import BaseWSGIServer, make_server


RESOURCE_DIR = Path(getattr(sys, "_MEIPASS", Path(__file__).resolve().parent))
if getattr(sys, "frozen", False):
    APP_DATA_DIR = Path(
        os.environ.get("LOCALAPPDATA", Path.home() / "AppData" / "Local")
    ) / "CartasYCorreos"
    APP_DATA_DIR.mkdir(parents=True, exist_ok=True)
    DB_PATH = APP_DATA_DIR / "programa.sqlite3"
    bundled_database = RESOURCE_DIR / "programa.sqlite3"
    if not DB_PATH.exists() and bundled_database.exists():
        shutil.copy2(bundled_database, DB_PATH)
else:
    DB_PATH = RESOURCE_DIR / "programa.sqlite3"
SENDER_EMAIL = "croyamid@hotmail.com"
MAX_UPLOAD_BYTES = 15 * 1024 * 1024
EMAIL_PATTERN = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")
HTTP_SERVER: BaseWSGIServer | None = None

app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = MAX_UPLOAD_BYTES


@contextmanager
def db_connection() -> Iterator[sqlite3.Connection]:
    connection = sqlite3.connect(DB_PATH)
    connection.row_factory = sqlite3.Row
    try:
        yield connection
        connection.commit()
    except Exception:
        connection.rollback()
        raise
    finally:
        connection.close()


def initialize_database() -> None:
    with db_connection() as connection:
        connection.execute(
            """CREATE TABLE IF NOT EXISTS settings (
                key TEXT PRIMARY KEY,
                value BLOB NOT NULL
            )"""
        )
        connection.execute(
            """CREATE TABLE IF NOT EXISTS contacts (
                id INTEGER PRIMARY KEY,
                company TEXT NOT NULL,
                email TEXT NOT NULL,
                status TEXT NOT NULL DEFAULT 'Pendiente',
                last_error TEXT,
                sent_at TEXT
            )"""
        )
        connection.execute("DELETE FROM settings WHERE key = 'client_id'")
        connection.execute(
            """UPDATE contacts SET status = 'Pendiente', last_error = NULL
               WHERE status = 'Error'"""
        )


def get_setting(key: str) -> bytes | None:
    with db_connection() as connection:
        row = connection.execute(
            "SELECT value FROM settings WHERE key = ?", (key,)
        ).fetchone()
    return bytes(row["value"]) if row else None


def save_setting(key: str, value: bytes | str) -> None:
    encoded_value = value.encode("utf-8") if isinstance(value, str) else value
    with db_connection() as connection:
        connection.execute(
            "INSERT OR REPLACE INTO settings (key, value) VALUES (?, ?)",
            (key, encoded_value),
        )


def normalize_header(value: object) -> str:
    text = unicodedata.normalize("NFKD", str(value or "").strip().lower())
    return "".join(char for char in text if not unicodedata.combining(char))


def read_contacts(workbook_bytes: bytes) -> tuple[list[tuple[str, str]], list[str]]:
    workbook = openpyxl.load_workbook(
        io.BytesIO(workbook_bytes), read_only=True, data_only=True
    )
    worksheet = workbook.active
    header_row: int | None = None
    company_column: int | None = None
    email_column: int | None = None

    for row_number, row in enumerate(
        worksheet.iter_rows(min_row=1, max_row=min(worksheet.max_row, 20), values_only=True),
        start=1,
    ):
        headers = [normalize_header(value) for value in row]
        company_match = next(
            (index for index, value in enumerate(headers) if value in {"empresa", "compania", "company"}),
            None,
        )
        email_match = next(
            (index for index, value in enumerate(headers) if value in {"correo", "email", "e-mail"}),
            None,
        )
        if company_match is not None and email_match is not None:
            header_row = row_number
            company_column = company_match
            email_column = email_match
            break

    if header_row is None or company_column is None or email_column is None:
        raise ValueError(
            "No encontré encabezados reconocibles. La hoja debe tener columnas «empresa» y «correo»."
        )

    contacts: list[tuple[str, str]] = []
    errors: list[str] = []
    seen_emails: set[str] = set()
    for row_number, row in enumerate(
        worksheet.iter_rows(min_row=header_row + 1, values_only=True),
        start=header_row + 1,
    ):
        company = str(row[company_column] or "").strip() if company_column < len(row) else ""
        email = str(row[email_column] or "").strip() if email_column < len(row) else ""
        if not company and not email:
            continue
        if not company or not EMAIL_PATTERN.fullmatch(email):
            errors.append(f"Fila {row_number}: falta la empresa o el correo no es válido.")
            continue
        normalized_email = email.casefold()
        if normalized_email in seen_emails:
            errors.append(f"Fila {row_number}: el correo está duplicado y se omitió.")
            continue
        seen_emails.add(normalized_email)
        contacts.append((company, email))

    if not contacts:
        raise ValueError("No hay contactos válidos para importar.")
    return contacts, errors


def import_contacts(workbook_bytes: bytes) -> list[str]:
    contacts, errors = read_contacts(workbook_bytes)
    with db_connection() as connection:
        connection.execute("DELETE FROM contacts")
        connection.executemany(
            "INSERT INTO contacts (company, email) VALUES (?, ?)", contacts
        )
    return errors


def replace_span(paragraph: Any, start: int, end: int, replacement: str) -> Any:
    runs = paragraph.runs
    run_texts = [run.text for run in runs]
    if not runs or start < 0 or end > sum(map(len, run_texts)) or start >= end:
        raise ValueError("No se pudo sustituir el espacio reservado en la carta.")

    cursor = 0
    start_run = start_offset = end_run = end_offset = None
    for index, text in enumerate(run_texts):
        next_cursor = cursor + len(text)
        if start_run is None and cursor <= start < next_cursor:
            start_run, start_offset = index, start - cursor
        if cursor < end <= next_cursor:
            end_run, end_offset = index, end - cursor
            break
        cursor = next_cursor

    if start_run is None or end_run is None or start_offset is None or end_offset is None:
        raise ValueError("El marcador de empresa está dividido de forma no compatible.")
    if start_run == end_run:
        runs[start_run].text = (
            run_texts[start_run][:start_offset]
            + replacement
            + run_texts[start_run][end_offset:]
        )
        return runs[start_run]

    runs[start_run].text = run_texts[start_run][:start_offset] + replacement
    for index in range(start_run + 1, end_run):
        runs[index].text = ""
    runs[end_run].text = run_texts[end_run][end_offset:]
    return runs[start_run]


def format_company_name(run: Any) -> None:
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(12)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Arial")


def personalized_letter(template_bytes: bytes, company: str) -> bytes:
    document = Document(io.BytesIO(template_bytes))
    company_marker_found = False
    salutation_marker_found = False
    date_marker_found = False
    company_pattern = re.compile(r"Empresa:\s*(_+)", re.IGNORECASE)
    salutation_pattern = re.compile(r"Señores:\s*(_+)", re.IGNORECASE)

    for paragraph in document.paragraphs:
        text = paragraph.text
        matches = []
        if match := salutation_pattern.search(text):
            matches.append((match.start(1), match.end(1), "salutation"))
            salutation_marker_found = True
        if match := company_pattern.search(text):
            matches.append((match.start(1), match.end(1), "company"))
            company_marker_found = True
        for start, end, _ in sorted(matches, reverse=True):
            run = replace_span(paragraph, start, end, company.upper())
            format_company_name(run)
        if not date_marker_found:
            date_match = re.search(r"\[fecha\]", paragraph.text, flags=re.IGNORECASE)
            if date_match:
                replace_span(
                    paragraph,
                    date_match.start(),
                    date_match.end(),
                    date.today().strftime("%d/%m/%Y"),
                )
                date_marker_found = True

    if not company_marker_found:
        raise ValueError(
            "La carta no contiene el espacio esperado «Empresa: ____». Revisa la plantilla."
        )
    if not salutation_marker_found:
        raise ValueError(
            "La carta no contiene el espacio esperado «Señores: ____». Revisa la plantilla."
        )
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def list_contacts() -> list[dict[str, Any]]:
    with db_connection() as connection:
        rows = connection.execute(
            """SELECT id, company, email, status, last_error, sent_at
               FROM contacts ORDER BY id"""
        ).fetchall()
    return [dict(row) for row in rows]


def make_email_zip(contact_ids: list[int]) -> io.BytesIO:
    template_bytes = get_setting("template")
    if not template_bytes:
        raise ValueError("Primero carga la carta de Word.")
    all_contacts = list_contacts()
    contacts_by_id = {contact["id"]: contact for contact in all_contacts}
    missing_ids = set(contact_ids) - contacts_by_id.keys()
    if missing_ids:
        raise ValueError("La selección incluye empresas que ya no están en la lista.")
    contacts = [contacts_by_id[contact_id] for contact_id in contact_ids]
    if not contacts:
        raise ValueError("Selecciona al menos una empresa.")

    archive = io.BytesIO()
    with zipfile.ZipFile(archive, "w", zipfile.ZIP_DEFLATED) as output_zip:
        for contact in contacts:
            letter = personalized_letter(template_bytes, contact["company"])
            safe_company = re.sub(r'[<>:"/\\|?*]+', "-", contact["company"]).strip(" .")
            filename = f"{contact['id']:04d} - {safe_company or 'Empresa'}.docx"
            draft = EmailMessage(policy=SMTP)
            draft["From"] = SENDER_EMAIL
            draft["To"] = contact["email"]
            draft["Subject"] = (
                "Propuesta comercial de servicios publicitarios - Emisora Radio 33,3"
            )
            draft.set_content(
                f"Cordial saludo, {contact['company']}.\n\n"
                "Adjuntamos nuestra propuesta comercial personalizada. "
                "Agradecemos su atención.\n\n"
                "Atentamente,\nEmisora Radio 33,3"
            )
            draft.add_attachment(
                letter,
                maintype="application",
                subtype="vnd.openxmlformats-officedocument.wordprocessingml.document",
                filename=filename,
            )
            draft_name = f"{contact['id']:04d} - {safe_company or 'Empresa'}.eml"
            output_zip.writestr(f"Correos/{draft_name}", draft.as_bytes())
            output_zip.writestr(f"Cartas/{filename}", letter)

        output_zip.writestr(
            "LEEME.txt",
            (
                "Correos preparados: cada archivo .eml contiene su destinatario, "
                "asunto, texto y carta Word adjunta.\n"
                "Abre cada .eml con tu aplicación de correo, revisa el mensaje y "
                "pulsa Enviar. El ZIP no envía correos automáticamente.\n"
                "Si tu aplicación no abre .eml para editarlo, usa la carta .docx "
                "de la carpeta Cartas y copia el correo indicado en el mensaje .eml.\n"
            ),
        )
    archive.seek(0)
    return archive


@app.get("/")
def index():
    return send_file(RESOURCE_DIR / "index.html")


@app.get("/api/status")
def api_status():
    return jsonify(
        {
            "sender": SENDER_EMAIL,
            "template_loaded": get_setting("template") is not None,
            "contacts": list_contacts(),
            "desktop_app": getattr(sys, "frozen", False),
        }
    )


@app.post("/api/template")
def api_template():
    upload = request.files.get("file")
    if not upload or not upload.filename.lower().endswith(".docx"):
        return jsonify(error="Selecciona una carta de Word .docx."), 400
    content = upload.read()
    try:
        Document(io.BytesIO(content))
    except (PackageNotFoundError, zipfile.BadZipFile, ValueError) as exc:
        return jsonify(error=f"No se pudo leer el archivo Word: {exc}"), 400
    save_setting("template", content)
    return jsonify(message="Carta Word cargada.")


@app.post("/api/contacts")
def api_contacts():
    upload = request.files.get("file")
    if not upload or not upload.filename.lower().endswith(".xlsx"):
        return jsonify(error="Selecciona una hoja de Excel .xlsx."), 400
    try:
        errors = import_contacts(upload.read())
    except (
        ValueError,
        zipfile.BadZipFile,
        openpyxl.utils.exceptions.InvalidFileException,
        KeyError,
    ) as exc:
        return jsonify(error=str(exc)), 400
    return jsonify(
        message=f"Se importaron {len(list_contacts())} contacto(s).",
        warnings=errors,
    )


@app.post("/api/emails.zip")
def api_emails_zip():
    data = request.get_json(silent=True) or {}
    ids = data.get("ids")
    try:
        if not isinstance(ids, list) or not ids:
            raise ValueError("Selecciona al menos una empresa.")
        selected_ids = list(dict.fromkeys(int(value) for value in ids))
        archive = make_email_zip(selected_ids)
    except (ValueError, TypeError, zipfile.BadZipFile) as exc:
        return jsonify(error=str(exc)), 400
    return send_file(
        archive,
        mimetype="application/zip",
        as_attachment=True,
        download_name="correos_listos_para_enviar.zip",
    )


@app.post("/api/shutdown")
def api_shutdown():
    if HTTP_SERVER is None:
        return jsonify(error="El cierre desde la página solo está disponible en la aplicación instalada."), 409
    threading.Timer(0.5, HTTP_SERVER.shutdown).start()
    return jsonify(message="Programa cerrado. Ya puedes cerrar esta pestaña.")


@app.errorhandler(413)
def file_too_large(_error):
    return jsonify(error="El archivo supera el límite de 15 MB."), 413


initialize_database()


def start_desktop_app() -> None:
    global HTTP_SERVER
    try:
        HTTP_SERVER = make_server("127.0.0.1", 5000, app, threaded=True)
    except OSError as exc:
        ctypes_message = (
            f"No se pudo iniciar el programa en el puerto local 5000.\n\n"
            f"Cierra la otra instancia de Cartas y correos e inténtalo de nuevo.\n\n{exc}"
        )
        import ctypes

        ctypes.windll.user32.MessageBoxW(
            None, ctypes_message, "Cartas y correos", 0x10
        )
        return

    url = f"http://127.0.0.1:{HTTP_SERVER.server_port}"
    if os.environ.get("CARTAS_NO_BROWSER") != "1":
        threading.Timer(0.8, webbrowser.open, args=(url,)).start()
    try:
        HTTP_SERVER.serve_forever()
    finally:
        HTTP_SERVER.server_close()


if __name__ == "__main__":
    if getattr(sys, "frozen", False):
        start_desktop_app()
    else:
        app.run(host="127.0.0.1", port=5000, debug=False)
